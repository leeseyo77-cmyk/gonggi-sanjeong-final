"""report_excel — 공사기간 산정 부록 엑셀의 구조와 수식(수 초, 샘플 불필요).

수식 결과가 앱 값과 같은지는 엑셀로 열어 대조했다(광주: 토목 2,198 / 건축 402 / 사업 전체 2,600 /
비작업 1,627 / 계 4,318일). 여기서는 수식이 앱의 계산 규칙대로 들어갔는지를 본다.
실행: python -m unittest discover -s tests -v
"""
import io
import unittest

import openpyxl

import _support  # noqa: F401  (저장소 경로 등록)
import report_excel as rx

DATA = {
    "title": "테스트 사업",
    "start_date": "2026-10-07",
    "civil": {"combine": "max", "notes": [], "cats": [
        {"name": "토공", "lines": [{"name": "(공통)", "items": [
            {"name": "터파기", "spec": "토사", "qty": 1000.0, "unit": "㎥", "daily": 420.0,
             "rate_unit": "㎥/일", "crew": 1, "days": 3, "source": "표준품셈 조기준(굴착기 1대/조)", "basis": ""}]}]},
        {"name": "구조물공", "lines": [
            {"name": "A동", "items": [
                {"name": "거푸집", "spec": "", "qty": 500.0, "unit": "㎡", "daily": 25.0, "rate_unit": "㎡/일",
                 "crew": 2, "days": 10, "source": "노무비역산(형틀목공)", "basis": "제1호표 …"}]},
            {"name": "B동", "items": [
                {"name": "거푸집", "spec": "", "qty": 200.0, "unit": "㎡", "daily": 25.0, "rate_unit": "㎡/일",
                 "crew": 1, "days": 8, "source": "노무비역산(형틀목공)", "basis": ""}]}]}]},
    "discs": [
        {"name": "건축", "use": True, "mode": "unit_parallel", "units": [
            {"name": "관리동", "packages": [{"group": "철근콘크리트공사", "lead": "형틀목공", "crew": 2, "days": 13,
                                          "omitted": 1, "items": [
                {"name": "유로폼", "spec": "", "unit": "㎡", "qty": 200.0, "rate": 0.13, "note": ""}]}]}]},
        {"name": "전기", "use": True, "mode": "unit_parallel", "units": []},
        {"name": "기계", "use": False, "mode": "unit_parallel", "units": []},
    ],
    "elec_after_arch": True,
    "nonwork": {"station": "강원도 강릉", "min_rest": 8,
                "conditions": [("혹서기(33℃ 이상)", [0] * 6 + [7.4, 6.0] + [0] * 4)],
                "rows": [{"월": "2026-12", "대상일수": 25, "달력일수": 31, "A": 3.63, "B": 4},
                         {"월": "2027-01", "대상일수": 31, "달력일수": 31, "A": 5.8, "B": 6}]},
    "holiday_years": {2026: [5] * 12, 2027: [6] * 12},
    "periods": {"prep": 61, "wrapup": 30, "commission": 0, "prep_label": "상수도공사"},
}


class AppendixTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(io.BytesIO(rx.build_appendix_xlsx(DATA)))

    def test_sheets_in_order(self):
        self.assertEqual(self.wb.sheetnames, [
            "총괄", "부록1. 작업일수 산정근거(토목)", "부록2. 작업일수 산정근거(건축)",
            "부록3. 작업일수 산정근거(전기 및 계측제어)", "부록4. 작업일수 산정근거(기계)",
            "별표1. 비작업일수 산정", "별표2. 기상조건별 비작업일수", "별표3. 법정공휴일수", "별표4. 준비기간"])

    def test_civil_formulas_follow_app_rules(self):
        ws = self.wb["부록1. 작업일수 산정근거(토목)"]
        formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)
                    and c.value.startswith("=")]
        self.assertIn("=ROUNDUP(D6/(F6*H6),0)", formulas)              # 항목 = ROUNDUP(수량÷(작업량×조))
        self.assertTrue(any(f.startswith("=MAX(I") for f in formulas))  # 구간 병행 → 최댓값
        self.assertEqual(ws["I4"].value, "=MAX(I5,I7)")                 # 토목 = 대공종 최댓값(최장)

    def test_discipline_bottleneck_formula(self):
        ws = self.wb["부록2. 작업일수 산정근거(건축)"]
        vals = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        self.assertTrue(any(v.startswith("=ROUNDUP(ROUND(G") for v in vals))
        self.assertTrue(any("다른 직종만 쓰는 항목 1개" in v for v in vals))

    def test_summary_chain_excludes_unreflected(self):
        ws = self.wb["총괄"]
        f = next(c.value for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.startswith("='부록1"))
        work = next(ws.cell(r, 2).value for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "사업 전체")
        # 토목 + MAX(건축, 건축 완료 후 전기); 미반영 기계는 빠진다
        self.assertIn("부록2", work)
        self.assertIn("부록3", work)
        self.assertNotIn("부록4", work)
        self.assertTrue(f.startswith("='부록1"))

    def test_nonwork_monthly_formulas(self):
        ws = self.wb["별표1. 비작업일수 산정"]
        vals = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        self.assertTrue(any(v.startswith("=ROUND(M7*M8/M6,0)") for v in vals))   # 중복 = A×B÷일수
        self.assertTrue(any(v.startswith("=ROUND(8*M6/31,0)") for v in vals))    # 부분 월 최소 일수 안분
        self.assertTrue(any(v.startswith("=MAX(") for v in vals))

    def test_work_formula_helper(self):
        f = rx._work_formula("T", {"건축": "A", "기계": "M", "전기": "E"}, elec_after_arch=False)
        self.assertEqual(f, "=T+MAX(A,M,M+E)")
        f = rx._work_formula("T", {"건축": "A", "전기": "E"}, elec_after_arch=True)
        self.assertEqual(f, "=T+MAX(A,A+E)")
        self.assertEqual(rx._work_formula("T", {}, True), "=T")


if __name__ == "__main__":
    unittest.main()
