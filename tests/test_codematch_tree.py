"""코드매칭형 변형(공주정수장형) — 번호 트리 공종 분리, 일위대가 열 순서, 품셈 작업조 환산,
일위대가에 적힌 일당 시공량, 산근 코드 따라가기(샘플 불필요).

실행: python -m unittest discover -s tests -v
"""
import unittest

import openpyxl

import _support  # noqa: F401  (저장소 경로 등록)
import crew_space as cs
import universal_parser as up

CM = up.TEMPLATES[1]          # 코드매칭형(내역서산근)


def _row(n=20, **cells):
    r = [None] * n
    for k, v in cells.items():
        r[int(k[1:])] = v
    return r


def _items_sheet(wb):
    ws = wb.create_sheet("내역서")
    rows = [
        _row(c0="내역서"), _row(c0="공 종 명", c1="규 격", c2="수량", c3="단위"), _row(c4="단 가"),
        _row(c0="가. 직접공사비", c2=1, c3="sum"),
        _row(c0="Ⅱ. 상수도시설", c2=1, c3="sum"),
        _row(c0=" A. UV 접촉조", c2=1, c3="sum"),
        _row(c0=" A.1 UV 접촉조"),
        _row(c0="  1 공통공사"), _row(c0="   1.1 가설공사"),
        _row(c0="    1.1.1 시스템비계", c1="10m이하", c2=1003, c3="㎡", c12="No.3", c17="S1"),
        _row(c0="  2 현장콘크리트타설공사"), _row(c0="   2.2 거푸집"),
        _row(c0="    2.2.4 유로폼", c1="복잡", c2=1197, c3="㎡", c12="No.9", c17="E1"),
        _row(c0=" B. 활성탄 흡착지", c2=1, c3="sum"),
        _row(c0="  2 현장콘크리트타설공사"),
        _row(c0="    2.2.1 유로폼", c1="복잡", c2=5768, c3="㎡", c12="No.9", c17="E1"),
        _row(c0="나. 안전시설", c2=1, c3="sum"),
        _row(c0=" Ⅰ 안전관리계획서 및 검토"),
        _row(c0="  안전관리계획서 작성비", c2=1, c3="식", c12="No.200", c17="Z1"),
    ]
    for r in rows:
        ws.append(r)


class TreeParseTest(unittest.TestCase):
    def test_numbered_tree(self):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        _items_sheet(wb)
        items = up.parse_items_generic(wb, CM)
        got = [(i["name"], i["category"], i["line"]) for i in items]
        self.assertEqual(got, [
            ("시스템비계", "공통공사", "UV 접촉조"),
            ("유로폼", "현장콘크리트타설공사", "UV 접촉조"),
            ("유로폼", "현장콘크리트타설공사", "활성탄 흡착지"),         # 구조물이 구간
            ("안전관리계획서 작성비", "안전시설 (공사비 외)", ""),        # 직접공사비 밖
        ])
        self.assertEqual(items[0]["code"], "S1")


def _labor_book(swapped=True):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    w = wb.create_sheet("노임")
    for nm, rate in (("보통인부", 172698), ("형틀목공", 275379), ("철근공", 268932), ("비계공", 284440)):
        w.append([nm, "일반공사 직종", "인", rate])
    ho = wb.create_sheet("일위대가_호표")

    def blk(spec, name, qty=None, unit=None, labor=None, code=None, comp=None):
        a, b = (spec, name) if swapped else (name, spec)
        return _row(c0=a, c1=b, c2=qty, c3=unit, c9=labor, c14=code, c17=comp)

    ho.append(_row(c0="일위대가_호표"))
    ho.append(_row(c0="규 격", c1="공 종 명") if swapped else _row(c0="공 종 명", c1="규 격"))
    for r in [
        # 표준시장단가: 노무비만 — 품셈 작업조(형틀목공 4 + 보통인부 1)로 환산
        blk("복잡, 0～7m", "No.9  유로폼", 1, "㎡", 43418, code="E1"),
        blk("복잡, 0～7m", "유로폼", 1, "㎡", 43418, comp="GCED402030000"),
        # 일위대가에 적힌 일당 시공량
        blk("토사, 기계", "No.148  절토면 고르기", 1, "M2", 1227, code="J1"),
        blk("1.0 M3", "굴삭기(무한궤도)", 8, "HR", 472200, comp="E000"),
        blk("일당 시공량 390㎡", "적용금액(1/390)", 1, "식", 1227),
        # 산근 코드를 가리키는 구성 행(설치 + 해체) — 단계별 일수 합
        blk("L=5m이하", "No.118  사보강재 설치및철거", 1, "본", 255354, code="X1"),
        blk("L=5m이하", "H-Beam 설치(버팀)", 1, "본", 165256, comp="HK1"),
        blk("L=5m이하", "H-Beam 해체(버팀)", 1, "본", 90098, comp="HK2"),
    ]:
        ho.append(r)
    return wb


class LaborVariantTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.res, cls.tr = up.parse_labor_derived_by_template(
            _labor_book(), CM, with_trace=True, code_daily={"HK1": (12.0, "본"), "HK2": (16.0, "본")})

    def test_market_price_block_uses_pumsem_crew(self):
        # (형틀목공 4 × 275,379 + 보통인부 1 × 172,698) ÷ 43,418 = 29.35㎡/조 (1인 기준이면 4.0)
        self.assertAlmostEqual(self.res["E1"][0], (4 * 275379 + 172698) / 43418, places=2)
        self.assertEqual(self.tr["E1"]["방식"], "노무비 ÷ 작업조 노임(품셈)")

    def test_stated_daily_output(self):
        self.assertEqual(self.res["J1"][:2], (390.0, "M2"))
        self.assertEqual(self.res["J1"][2], "일당 시공량")

    def test_codematch_sangeun_traversal(self):
        self.assertAlmostEqual(self.res["X1"][0], 1 / (1 / 12 + 1 / 16), places=3)
        self.assertEqual(self.res["X1"][2], "산근 A·Q")

    def test_column_order_detected(self):
        # 열 순서가 화성형('공종명 | 규격')이어도 같은 결과
        res2 = up.parse_labor_derived_by_template(_labor_book(swapped=False), CM)
        self.assertAlmostEqual(res2["E1"][0], self.res["E1"][0], places=4)


class CrewStepsTest(unittest.TestCase):
    def test_steps(self):
        self.assertEqual(cs.crew_steps("유로폼", "복잡"), [({"형틀목공": 4, "보통인부": 1}, 35)])
        self.assertEqual(len(cs.crew_steps("시스템 동바리 설치 및 해체")), 2)     # 설치 → 해체
        self.assertEqual(len(cs.crew_steps("시스템 동바리 설치")), 1)
        # 공장가공이면 현장은 조립만, 현장가공 및 조립이면 가공 + 조립
        self.assertEqual(cs.crew_steps("철근가공조립", "철근 공장가공및조립,Type-Ⅱ")[0][0]["철근공"], 5)
        self.assertEqual(len(cs.crew_steps("철근 현장가공 및 조립", "Type-Ⅱ")), 2)
        self.assertEqual(cs.crew_steps("철근 현장가공 및 조립", "TYPE-Ⅰ")[1][1], 3.4)
        self.assertIsNone(cs.crew_steps("철근콘크리트 깨기"))
        self.assertIsNone(cs.crew_steps("잡철물 제작 및 설치"))
        self.assertEqual(cs.crew_steps("철근콘크리트 타설/펌프차", "TYPE-2")[0][0]["콘크리트공"], 4)


if __name__ == "__main__":
    unittest.main()
