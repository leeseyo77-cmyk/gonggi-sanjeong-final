"""앱 스모크·회귀 테스트 — 광주 제1정수장 샘플로 app.py를 헤드리스 실행한다(수 분).

샘플이 없거나 GONGGI_SKIP_APP=1이면 건너뛴다.
실행: python -m unittest discover -s tests -v

기준값(2,198일·3,759일 등)은 현재 산정 모델의 결과다. 모델을 의도적으로 바꿨다면
결과를 확인한 뒤 여기 숫자를 갱신한다.
"""
import os
import unittest
import zipfile

import openpyxl

from _support import (GWANGJU, OUT, app_test, button, exceptions, have, metric,
                      widget)

SKIP = os.environ.get("GONGGI_SKIP_APP") == "1" or not have(*GWANGJU.values())
WHY = "광주 샘플이 없거나 GONGGI_SKIP_APP=1"


@unittest.skipIf(SKIP, WHY)
class CivilTest(unittest.TestCase):
    def test_civil_work_days_baseline(self):
        at = app_test(civil=GWANGJU["토목"])
        self.assertEqual(exceptions(at), [])
        # 노무비 역산은 병목 직종 기준(1조 = 직종별 1인). 1인 합산 기준이던 때는 2,517일
        self.assertEqual(at.session_state["total_work_days"], 2198)


@unittest.skipIf(SKIP, WHY)
class ProjectTabTest(unittest.TestCase):
    def test_defaults_then_include_electrical(self):
        at = app_test(civil=GWANGJU["토목"],
                      discs={"건축": GWANGJU["건축"], "전기": GWANGJU["전기"]})
        self.assertEqual(exceptions(at), [])
        # 기본: 토목 + 건축(동 병행) — 전기는 공기 미반영
        self.assertEqual(metric(at, "사업 전체 순작업일수"), "3759일")

        # 전기를 켜면 건축 마감 후 착수(기본) → 2,198 + 1,561 + 984
        widget(at, "checkbox", "disc_use_전기").set_value(True).run()
        self.assertEqual(exceptions(at), [])
        self.assertEqual(metric(at, "사업 전체 순작업일수"), "4743일")

        # 비작업일수 계산기로 넘기기
        button(at, "📥 전체 순작업일수").click().run()
        self.assertEqual(exceptions(at), [])
        self.assertEqual(at.session_state["weather_work_days"], 4743)

    def test_crew_allocation_goes_to_larger_building(self):
        at = app_test(civil=GWANGJU["토목"], discs={"건축": GWANGJU["건축"]})
        widget(at, "number_input", "disc_tc_건축").set_value(8).run()
        self.assertEqual(exceptions(at), [])
        # 면 작업 상한 4조 → 활성탄흡착지 4조가 한계, 402일
        self.assertEqual(metric(at, "건축 작업일수"), "402일")


@unittest.skipIf(SKIP, WHY)
class CrewReverseAndScheduleTest(unittest.TestCase):
    """비작업일수 계산 후에만 나오는 토목 조수 역산 화면과 예정공정표 엑셀."""

    @classmethod
    def setUpClass(cls):
        cls.at = app_test(civil=GWANGJU["토목"], discs={"건축": GWANGJU["건축"]})
        widget(cls.at, "number_input", "disc_tc_건축").set_value(8).run()
        button(cls.at, "📊 비작업일수 계산").click().run()

    def test_project_settings_survive_weather_calc(self):
        # '비작업일수 계산'이 탭 중간에서 바로 재실행하면 사업 전체 탭 위젯이 그려지지 않아
        # Streamlit이 그 상태(조수·업로드)를 지웠다 → 재실행을 실행 끝으로 미뤄 고쳤다
        self.assertEqual(self.at.session_state["disc_tc_건축"], 8)
        self.assertEqual(metric(self.at, "건축 작업일수"), "402일")

    def test_appendix_workbook_generated(self):
        at = self.at
        for f in OUT.glob("공사기간 산정 부록_*.xlsx"):
            f.unlink()
        button(at, "📥 부록 엑셀 생성").click().run()
        self.assertEqual(exceptions(at), [])
        # 부록이 다시 묶은 토목 작업일수가 앱 산정값과 같아야 한다(다르면 경고가 뜬다)
        self.assertFalse(any("부록의 토목 작업일수" in w.value for w in at.warning))
        files = list(OUT.glob("공사기간 산정 부록_*.xlsx"))
        self.assertEqual(len(files), 1)
        wb = openpyxl.load_workbook(files[0])
        self.assertEqual(wb.sheetnames[:3], ["총괄", "부록1. 작업일수 산정근거(토목)",
                                             "부록2. 작업일수 산정근거(건축)"])
        self.assertIn("별표1. 비작업일수 산정", wb.sheetnames)

    def test_unreachable_target_is_flagged(self):
        at = self.at
        widget(at, "number_input", "target_months_input").set_value(24.0).run()
        self.assertEqual(exceptions(at), [])
        self.assertIsNotNone(metric(at, "역산안 (필요 조수)"))
        self.assertTrue(any("목표 공기가 무리" in e.value for e in at.error))

        widget(at, "number_input", "target_months_input").set_value(240.0).run()
        self.assertFalse(any("목표 공기가 무리" in e.value for e in at.error))

    def test_schedule_excel_states_crew_assumption(self):
        at = self.at
        for f in OUT.glob("예정공정표_*.xlsx"):
            f.unlink()
        button(at, "📥 예정공정표 엑셀 생성").click().run()
        self.assertEqual(exceptions(at), [])
        files = list(OUT.glob("예정공정표_*.xlsx"))
        self.assertEqual(len(files), 1)
        ws = openpyxl.load_workbook(files[0]).active
        notes = [c.value for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.startswith("※")]
        self.assertTrue(any("가정값" in n for n in notes), notes)
        # 사업 전체 공기 탭의 건축이 동 단위로 공정표에 들어간다
        col_a = [ws.cell(r, 1).value for r in range(5, ws.max_row + 1)]
        col_b = [ws.cell(r, 2).value for r in range(5, ws.max_row + 1)]
        self.assertIn("건축공사", col_a)
        self.assertIn("활성탄흡착지", col_b)
        with zipfile.ZipFile(files[0]) as z:
            self.assertIn("xl/drawings/drawing1.xml", z.namelist())


if __name__ == "__main__":
    unittest.main()
