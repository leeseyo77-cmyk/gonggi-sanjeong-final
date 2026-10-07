# -*- coding: utf-8 -*-
"""
report_excel.py — 공사기간 산정 부록 엑셀

발주처 제출용 '공사기간 산정' 부록(사용자 샘플: 영덕 영해 하수처리시설)의 구성을 따른다.
  총괄                     작업일수·비작업일수·준비·정리·시운전 → 공사기간(일·개월)
  부록n 작업일수 산정근거   토목: 대공종 → 라인(구간·구조물) → 항목
                           건축 등: 동(설비) → 공종 → 항목(병목 직종 작업량)
  별표1 비작업일수 산정     연도별 월 표(고시 제8조: 공휴일 + 기후여건 − 중복, 월 최소 8일)
  별표2 기상조건별 비작업일수, 별표3 법정공휴일수, 별표4 준비기간

작업일수·소계·합계·중복일수·비작업일수는 수식으로 넣는다. 검토자가 조수나 1일 작업량을 고치면
엑셀이 다시 계산한다(샘플과 같다). 수식 결과가 앱 산정값과 같도록 앱의 계산 방식을 그대로 옮겼다:
  토목  항목 = ROUNDUP(수량 ÷ (1일 작업량 × 조)), 라인 = 항목 합(라인 안은 순차),
        대공종 = 라인 중 최댓값(라인끼리 병행), 토목 = 대공종 최댓값(최장) 또는 합(합산)
  분야  공종 = ROUNDUP(병목 직종 작업량 ÷ 조), 동 = 공종 합, 분야 = 동 최댓값(병행) 또는 합(순차)
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "맑은 고딕"
_thin = Side(style="thin", color="808080")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_HDR = PatternFill("solid", fgColor="D9E1F2")
_CAT = PatternFill("solid", fgColor="F2F2F2")
_SUB = PatternFill("solid", fgColor="FFF2CC")

# 적정 공사기간 확보를 위한 가이드라인(국토교통부) — 공사 유형별 준비기간(일)
PREP_GUIDE = [
    ("공동주택", 45), ("고속도로공사", 180), ("철도공사", 90), ("포장공사(신설)", 50),
    ("포장공사(수선)", 60), ("공동구공사", 80), ("상수도공사", 60), ("하천공사", 40),
    ("항만공사", 40), ("강교가설공사", 90), ("PC교량 공사", 70), ("교량보수공사", 60),
]

DISC_LABEL = {"건축": "건축", "건축기계설비": "건축기계설비", "기계": "기계",
              "전기": "전기 및 계측제어", "조경": "조경"}


def _qty_fmt(v) -> str:
    """정수는 '76', 소수는 '76.5'처럼(‘#,##0.###’은 정수 뒤에 점이 남는다)."""
    try:
        return "#,##0" if float(v).is_integer() else "#,##0.0##"
    except (TypeError, ValueError):
        return "General"


def _text_lines(text: str, width_chars: float) -> int:
    """열 너비(문자 수)에 줄바꿈될 때의 대략 줄 수 — 한글·전각은 2칸으로 센다."""
    import math
    w = sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(text or ""))
    return max(1, math.ceil(w / max(1.0, width_chars * 1.15)))


def _q(sheet: str) -> str:
    """수식용 시트 이름(작은따옴표로 감싼다)."""
    return "'" + sheet.replace("'", "''") + "'"


class _Sheet:
    """셀 쓰기 도우미: 글꼴·정렬·테두리·채우기를 한 번에."""

    def __init__(self, ws):
        self.ws = ws

    def put(self, r, c, v=None, *, bold=False, size=9, h="center", fill=None, fmt=None,
            border=True, wrap=False, color="000000"):
        x = self.ws.cell(row=r, column=c, value=v)
        x.font = Font(name=FONT, size=size, bold=bold, color=color)
        x.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
        if border:
            x.border = _BORDER
        if fill:
            x.fill = fill
        if fmt:
            x.number_format = fmt
        return x

    def row(self, r, values, **kw):
        for c, v in enumerate(values, start=1):
            self.put(r, c, v, **kw)

    def widths(self, ws_widths: Sequence[float]):
        for i, w in enumerate(ws_widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = w

    def title(self, text, ncols):
        self.ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        self.put(1, 1, text, bold=True, size=13, h="left", border=False)
        self.ws.row_dimensions[1].height = 24


def _print_setup(ws, landscape=False, title_rows: Optional[str] = None):
    ws.page_setup.paperSize = 9                     # A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if title_rows:
        ws.print_title_rows = title_rows


# ── 부록: 토목 ───────────────────────────────────────────────────────────
def _civil_sheet(wb, name: str, civil: Dict[str, Any]) -> str:
    """토목 작업일수 산정근거. 반환: 토목 합계 셀 주소(시트!셀)."""
    ws = wb.create_sheet(name)
    S = _Sheet(ws)
    hdr = ["번호", "공 종 명", "규 격", "수량", "단위", "1일 작업량", "작업량 단위", "조",
           "작업일수(일)", "출처", "산출근거"]
    S.widths([6, 34, 26, 10, 6, 11, 10, 5, 10, 18, 58])
    S.title("◈ " + name.split(". ", 1)[-1], len(hdr))
    S.put(2, 1, "(1) 토목공사 — 라인(구간·구조물) 안은 순차로 더하고, 라인끼리는 병행이라 가장 긴 "
                "라인이 대공종 작업일수입니다.", h="left", border=False)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(hdr))
    S.row(3, hdr, bold=True, fill=_HDR, wrap=True)
    ws.row_dimensions[3].height = 30
    total_r = 4
    S.put(total_r, 1, "", fill=_CAT)
    S.put(total_r, 2, "토목공사 계", bold=True, h="left", fill=_CAT)
    for c in range(3, len(hdr) + 1):
        S.put(total_r, c, None, fill=_CAT)
    combine_sum = civil.get("combine") == "sum"
    S.put(total_r, 11, "대공종 " + ("합(합산·순차)" if combine_sum else "최댓값(최장·병행)"),
          h="left", fill=_CAT)

    r = total_r + 1
    cat_cells = []
    for ci, cat in enumerate(civil.get("cats", []), start=1):
        cat_r = r
        S.put(cat_r, 1, ci, bold=True, fill=_CAT)
        S.put(cat_r, 2, cat["name"], bold=True, h="left", fill=_CAT)
        for c in range(3, len(hdr) + 1):
            S.put(cat_r, c, None, fill=_CAT)
        r += 1
        lines = [ln for ln in cat["lines"] if ln["items"]]
        line_cells = []
        multi = len(lines) > 1
        for ln in lines:
            if multi:
                sub_r = r
                S.put(sub_r, 2, f"▸ {ln['name']}", h="left", fill=_SUB)
                for c in (1, *range(3, len(hdr) + 1)):
                    S.put(sub_r, c, None, fill=_SUB)
                S.put(sub_r, 11, "구간 소계(구간 안 항목 순차)", h="left", fill=_SUB)
                r += 1
            first = r
            for it in ln["items"]:
                S.put(r, 1, None)
                S.put(r, 2, it["name"], h="left")
                S.put(r, 3, it.get("spec", ""), h="left")
                S.put(r, 4, it["qty"], fmt=_qty_fmt(it["qty"]))
                S.put(r, 5, it.get("unit", ""))
                S.put(r, 6, it["daily"], fmt="#,##0.0##")
                S.put(r, 7, it.get("rate_unit", ""))
                S.put(r, 8, it["crew"])
                S.put(r, 9, f"=ROUNDUP(D{r}/(F{r}*H{r}),0)", fmt="#,##0")
                S.put(r, 10, it.get("source", ""), h="left", wrap=True)
                S.put(r, 11, it.get("basis", ""), h="left", wrap=True)
                _nl = max(_text_lines(it.get("basis", ""), 58), _text_lines(it.get("spec", ""), 26))
                if _nl > 1:
                    ws.row_dimensions[r].height = 12.5 * _nl + 2
                r += 1
            last = r - 1
            if multi:
                S.put(sub_r, 9, f"=SUM(I{first}:I{last})", bold=True, fmt="#,##0", fill=_SUB)
                line_cells.append(f"I{sub_r}")
            else:
                line_cells.append(f"SUM(I{first}:I{last})")
        if multi:
            S.put(cat_r, 9, "=MAX(" + ",".join(line_cells) + ")", bold=True, fmt="#,##0", fill=_CAT)
            S.put(cat_r, 11, f"구간 {len(lines)}개 병행 → 가장 긴 구간", h="left", fill=_CAT)
        else:
            S.put(cat_r, 9, "=" + (line_cells[0] if line_cells else "0"), bold=True, fmt="#,##0",
                  fill=_CAT)
        cat_cells.append(f"I{cat_r}")
    fn = "SUM" if combine_sum else "MAX"
    S.put(total_r, 9, f"={fn}(" + ",".join(cat_cells) + ")" if cat_cells else 0, bold=True,
          fmt="#,##0", fill=_CAT)
    for note in civil.get("notes", []):
        r += 1
        S.put(r, 1, note, h="left", border=False, size=8)
    ws.freeze_panes = "A4"
    _print_setup(ws, landscape=True, title_rows="3:3")
    return f"{_q(name)}!I{total_r}"


# ── 부록: 건축·기계·전기·조경 ─────────────────────────────────────────────
def _disc_sheet(wb, name: str, disc: Dict[str, Any]) -> str:
    """분야 작업일수 산정근거. 반환: 분야 합계 셀 주소."""
    ws = wb.create_sheet(name)
    S = _Sheet(ws)
    hdr = ["번호", "공 종 명", "규 격", "단위", "수량", "주 직종 인수\n(인/단위)", "작업량\n(인·일)",
           "조", "작업일수(일)", "비 고"]
    S.widths([6, 34, 28, 6, 10, 12, 11, 5, 10, 40])
    label = DISC_LABEL.get(disc["name"], disc["name"])
    S.title("◈ " + name.split(". ", 1)[-1], len(hdr))
    parallel = disc.get("mode") != "seq"
    S.put(2, 1, f"({label}공사) 동·설비 안의 공종은 순차, " +
                ("동·설비끼리는 병행이라 가장 긴 동이 작업일수입니다." if parallel else
                 "동·설비도 순차로 짓습니다.") +
                " 공종 작업일수 = 가장 많이 드는 직종(병목)의 작업량 ÷ 조 (1조 = 직종별 1인)."
                + ("" if disc.get("use", True) else "  ※ 공기 미반영(참고)"),
          h="left", border=False)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(hdr))
    S.row(3, hdr, bold=True, fill=_HDR, wrap=True)
    ws.row_dimensions[3].height = 30
    total_r = 4
    S.put(total_r, 1, "", fill=_CAT)
    S.put(total_r, 2, f"{label}공사 계", bold=True, h="left", fill=_CAT)
    for c in range(3, len(hdr) + 1):
        S.put(total_r, c, None, fill=_CAT)
    S.put(total_r, 10, "동·설비 " + ("최댓값(병행)" if parallel else "합(순차)"), h="left", fill=_CAT)
    r = total_r + 1
    unit_cells = []
    for ui, unit in enumerate(disc.get("units", []), start=1):
        unit_r = r
        S.put(unit_r, 1, ui, bold=True, fill=_CAT)
        S.put(unit_r, 2, unit["name"], bold=True, h="left", fill=_CAT)
        for c in range(3, len(hdr) + 1):
            S.put(unit_r, c, None, fill=_CAT)
        S.put(unit_r, 10, "공종 순차 합", h="left", fill=_CAT)
        r += 1
        pkg_cells = []
        for pk in unit["packages"]:
            pk_r = r
            S.put(pk_r, 2, f"▸ {pk['group']}", h="left", fill=_SUB)
            _lead = "노무비 역산(직종 정보 없음)" if pk["lead"] == "노무비역산" else pk["lead"]
            S.put(pk_r, 3, f"주 직종: {_lead}", h="left", fill=_SUB)
            for c in (1, 4, 5, 6):
                S.put(pk_r, c, None, fill=_SUB)
            S.put(pk_r, 8, pk["crew"], fill=_SUB)
            r += 1
            first = r
            for it in pk["items"]:
                S.put(r, 1, None)
                S.put(r, 2, it["name"], h="left")
                S.put(r, 3, it.get("spec", ""), h="left")
                S.put(r, 4, it.get("unit", ""))
                S.put(r, 5, it["qty"], fmt=_qty_fmt(it["qty"]))
                S.put(r, 6, it["rate"], fmt="0.0#####")
                S.put(r, 7, f"=E{r}*F{r}", fmt="#,##0.0#")
                S.put(r, 8, None)
                S.put(r, 9, None)
                S.put(r, 10, it.get("note", ""), h="left")
                r += 1
            last = r - 1
            S.put(pk_r, 7, f"=SUM(G{first}:G{last})" if pk["items"] else 0, bold=True,
                  fmt="#,##0.0#", fill=_SUB)
            S.put(pk_r, 9, f"=ROUNDUP(ROUND(G{pk_r}/H{pk_r},6),0)", bold=True, fmt="#,##0", fill=_SUB)
            S.put(pk_r, 10, (f"다른 직종만 쓰는 항목 {pk['omitted']}개는 병목 작업량에 들어가지 않음"
                             if pk.get("omitted") else ""), h="left", fill=_SUB, size=8)
            pkg_cells.append(f"I{pk_r}")
        S.put(unit_r, 9, "=SUM(" + ",".join(pkg_cells) + ")" if pkg_cells else 0, bold=True,
              fmt="#,##0", fill=_CAT)
        unit_cells.append(f"I{unit_r}")
    fn = "MAX" if parallel else "SUM"
    S.put(total_r, 9, f"={fn}(" + ",".join(unit_cells) + ")" if unit_cells else 0, bold=True,
          fmt="#,##0", fill=_CAT)
    ws.freeze_panes = "A4"
    _print_setup(ws, landscape=True, title_rows="3:3")
    return f"{_q(name)}!I{total_r}"


# ── 별표: 비작업일수 ─────────────────────────────────────────────────────
def _nonwork_sheet(wb, name: str, nw: Dict[str, Any]) -> str:
    """연도별 월 표(샘플 '비작업일수 세부산정' 배치). 반환: 적용 비작업일수 합계 셀 주소."""
    ws = wb.create_sheet(name)
    S = _Sheet(ws)
    S.widths([18] + [7.5] * 12 + [9])
    S.title("▣ " + name.split(". ", 1)[-1], 14)
    mn = nw.get("min_rest", 8)
    S.put(2, 1, f"고시 제8조: 비작업일수 = 법정공휴일(A) + 기후여건(B) − 중복(C), "
                f"중복 = A × B ÷ 일수(반올림). 월 비작업일수가 {mn}일(주 40시간 근무제) 미만이면 {mn}일 적용."
                if mn else "고시 제8조: 비작업일수 = 법정공휴일(A) + 기후여건(B) − 중복(C), 중복 = A × B ÷ 일수(반올림).",
          h="left", border=False)
    ws.merge_cells("A2:N2")
    by_year: Dict[int, Dict[int, Dict[str, Any]]] = {}
    for row in nw.get("rows", []):
        y, m = map(int, str(row["월"]).split("-"))
        by_year.setdefault(y, {})[m] = row
    labels = ["일수", "법정공휴일(A)", "기후여건(B)", "월별중복일수(C)", "비작업일수",
              "주40시간 고려", "적용", "작업가능일수", "가동률(%)"]
    r = 4
    total_cells = []
    for y in sorted(by_year):
        months = by_year[y]
        S.put(r, 1, "구 분", bold=True, fill=_HDR)
        S.put(r, 2, f"{y}년", bold=True, fill=_HDR)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
        r += 1
        S.row(r, [""] + [f"{m}월" for m in range(1, 13)] + ["소계"], bold=True, fill=_HDR)
        r += 1
        base = r                      # base+0 일수, +1 A, +2 B, +3 C, +4 비작업, +5 최소, +6 적용, +7 작업, +8 가동률
        for i, lab in enumerate(labels):
            S.put(base + i, 1, lab, bold=(lab == "적용"), fill=_CAT if lab == "적용" else None)
        for m in range(1, 13):
            c = m + 1
            L = get_column_letter(c)
            mrow = months.get(m)
            if not mrow:
                for i in range(len(labels)):
                    S.put(base + i, c, None)
                continue
            S.put(base, c, mrow["대상일수"])
            S.put(base + 1, c, mrow["B"])
            S.put(base + 2, c, mrow["A"], fmt="0.0#")
            S.put(base + 3, c, f"=ROUND({L}{base+1}*{L}{base+2}/{L}{base},0)")
            S.put(base + 4, c, f"=ROUND({L}{base+1}+{L}{base+2}-{L}{base+3},0)")
            S.put(base + 5, c, f"=ROUND({mn}*{L}{base}/{mrow['달력일수']},0)" if mn else 0)
            S.put(base + 6, c, f"=MAX({L}{base+4},{L}{base+5})", bold=True, fill=_CAT)
            S.put(base + 7, c, f"={L}{base}-{L}{base+6}")
            S.put(base + 8, c, f"=ROUND({L}{base+7}/{L}{base}*100,1)", fmt="0.0")
        for i in range(len(labels)):
            rr = base + i
            if labels[i] == "가동률(%)":
                S.put(rr, 14, f"=ROUND(N{base+7}/N{base}*100,1)", fmt="0.0")
            else:
                S.put(rr, 14, f"=SUM(B{rr}:M{rr})", bold=(labels[i] == "적용"),
                      fmt="0.0#" if labels[i] == "기후여건(B)" else "#,##0",
                      fill=_CAT if labels[i] == "적용" else None)
        total_cells.append(f"N{base+6}")
        r = base + len(labels) + 1
    S.put(r, 1, "적용 비작업일수 계", bold=True, fill=_SUB)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
    S.put(r, 14, "=" + "+".join(total_cells) if total_cells else 0, bold=True, fmt="#,##0", fill=_SUB)
    total_ref = f"{_q(name)}!N{r}"
    r += 2
    S.put(r, 1, "※ 첫·마지막 달은 공사 구간에 든 일수만큼 공휴일·기후여건·최소 일수를 안분했습니다.",
          h="left", border=False, size=8)
    if nw.get("station"):
        r += 1
        S.put(r, 1, f"※ 기후여건: {nw['station']} 관측지점, 적정 공사기간 확보를 위한 가이드라인 부록 월평균",
              h="left", border=False, size=8)
    _print_setup(ws, landscape=True)
    return total_ref


def _conditions_sheet(wb, name: str, nw: Dict[str, Any]) -> None:
    ws = wb.create_sheet(name)
    S = _Sheet(ws)
    S.widths([30] + [7] * 12 + [9])
    S.title("▣ " + name.split(". ", 1)[-1], 14)
    S.put(2, 1, f"지점: {nw.get('station') or '-'}  (월평균 비작업일수, 가이드라인 부록)", h="left",
          border=False)
    ws.merge_cells("A2:N2")
    S.row(4, ["기상조건"] + [f"{m}월" for m in range(1, 13)] + ["계"], bold=True, fill=_HDR)
    r = 5
    first = r
    for label, vals in nw.get("conditions", []):
        S.put(r, 1, label, h="left")
        for m, v in enumerate(vals, start=2):
            S.put(r, m, v, fmt="0.0#")
        S.put(r, 14, f"=SUM(B{r}:M{r})", fmt="0.0#")
        r += 1
    if r > first:
        S.put(r, 1, "합계(기후여건)", bold=True, fill=_CAT)
        for c in range(2, 15):
            L = get_column_letter(c)
            S.put(r, c, f"=SUM({L}{first}:{L}{r-1})", bold=True, fmt="0.0#", fill=_CAT)
    else:
        S.put(r, 1, "선택한 기상조건이 없습니다.", h="left", border=False)
    _print_setup(ws, landscape=True)


def _holiday_sheet(wb, name: str, years: Dict[int, List[int]]) -> None:
    ws = wb.create_sheet(name)
    S = _Sheet(ws)
    S.widths([10] + [6] * 12 + [8])
    S.title("▣ " + name.split(". ", 1)[-1], 14)
    S.put(2, 1, "「관공서의 공휴일에 관한 규정」의 공휴일과 「근로기준법」의 근로자의 날(가이드라인 부록)",
          h="left", border=False)
    ws.merge_cells("A2:N2")
    S.row(4, ["구분"] + [f"{m}월" for m in range(1, 13)] + ["소계"], bold=True, fill=_HDR)
    r = 5
    for y in sorted(years):
        S.put(r, 1, f"{y}년")
        for m, v in enumerate(years[y], start=2):
            S.put(r, m, v)
        S.put(r, 14, f"=SUM(B{r}:M{r})", bold=True)
        r += 1
    _print_setup(ws, landscape=True)


def _prep_sheet(wb, name: str, prep_days: int, applied_label: str) -> None:
    ws = wb.create_sheet(name)
    S = _Sheet(ws)
    S.widths([22, 12, 4, 22, 12])
    S.title("▣ " + name.split(". ", 1)[-1], 5)
    S.put(2, 1, "적정 공사기간 확보를 위한 가이드라인(국토교통부) 공사 유형별 준비기간", h="left",
          border=False)
    ws.merge_cells("A2:E2")
    S.row(4, ["공 종", "준비기간(일)", None, "공 종", "준비기간(일)"], bold=True, fill=_HDR)
    half = (len(PREP_GUIDE) + 1) // 2
    for i in range(half):
        a = PREP_GUIDE[i]
        b = PREP_GUIDE[i + half] if i + half < len(PREP_GUIDE) else ("", None)
        S.put(5 + i, 1, a[0]); S.put(5 + i, 2, a[1])
        S.put(5 + i, 3, None, border=False)
        S.put(5 + i, 4, b[0]); S.put(5 + i, 5, b[1])
    r = 5 + half + 1
    S.row(r, ["적용", "준비기간(일)", None, "비 고", None], bold=True, fill=_HDR)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    S.put(r + 1, 1, applied_label or "-")
    S.put(r + 1, 2, prep_days, bold=True)
    S.put(r + 1, 3, None, border=False)
    S.put(r + 1, 4, "비작업일수 계산기의 준비기간 설정값", h="left")
    S.put(r + 1, 5, None)
    ws.merge_cells(start_row=r + 1, start_column=4, end_row=r + 1, end_column=5)
    _print_setup(ws)


# ── 총괄 ────────────────────────────────────────────────────────────────
def _work_formula(civil_ref: str, disc_refs: Dict[str, str], elec_after_arch: bool) -> str:
    """사업 전체 작업일수 = 토목 + 토목 이후 구간(사업 전체 공기 탭의 공정 연결과 같은 식)."""
    arch = [disc_refs[k] for k in ("건축", "건축기계설비") if k in disc_refs]
    mech, land, elec = disc_refs.get("기계"), disc_refs.get("조경"), disc_refs.get("전기")
    branches = list(arch)
    branches += [x for x in (mech, land) if x]
    if elec:
        pred = arch + ([mech] if mech else []) if elec_after_arch else ([mech] if mech else [])
        pred_s = ("MAX(" + ",".join(pred) + ")") if len(pred) > 1 else (pred[0] if pred else "0")
        branches.append(f"{pred_s}+{elec}")
    if not branches:
        return f"={civil_ref}"
    after = branches[0] if len(branches) == 1 else "MAX(" + ",".join(branches) + ")"
    return f"={civil_ref}+{after}"


def build_appendix_xlsx(data: Dict[str, Any]) -> bytes:
    """앱 산정 결과(data)로 공사기간 산정 부록 엑셀을 만든다(바이트)."""
    wb = Workbook()
    summary = wb.active
    summary.title = "총괄"
    # 부록 시트 이름은 번호를 차례로 붙인다(샘플의 '부록2' 중복 같은 일이 없게)
    civil_name = "부록1. 작업일수 산정근거(토목)"
    civil_ref = _civil_sheet(wb, civil_name, data["civil"])
    disc_refs: Dict[str, str] = {}
    disc_rows = []
    n = 2
    for d in data.get("discs", []):
        nm = f"부록{n}. 작업일수 산정근거({DISC_LABEL.get(d['name'], d['name'])})"
        ref = _disc_sheet(wb, nm, d)
        disc_rows.append((d, ref))
        if d.get("use", True):
            disc_refs[d["name"]] = ref
        n += 1
    nw_name = "별표1. 비작업일수 산정"
    nw_ref = _nonwork_sheet(wb, nw_name, data["nonwork"])
    _conditions_sheet(wb, "별표2. 기상조건별 비작업일수", data["nonwork"])
    _holiday_sheet(wb, "별표3. 법정공휴일수", data.get("holiday_years", {}))
    per = data["periods"]
    _prep_sheet(wb, "별표4. 준비기간", per["prep"], per.get("prep_label", ""))

    # 총괄
    S = _Sheet(summary)
    S.widths([24, 12, 12, 12, 12, 12, 12, 12])
    S.title(f"■ {data.get('title') or '공사'} 공사기간 산정 총괄", 8)
    S.row(3, ["구 분", "작업일수", "비작업일수", "준비기간", "정리기간", "시운전", "계(일)", "개월"],
          bold=True, fill=_HDR)
    r = 4
    S.put(r, 1, "토목", h="left")
    S.put(r, 2, f"={civil_ref}", fmt="#,##0")
    for c in range(3, 9):
        S.put(r, c, None)
    r += 1
    for d, ref in disc_rows:
        lab = DISC_LABEL.get(d["name"], d["name"]) + ("" if d.get("use", True) else " (공기 미반영·참고)")
        S.put(r, 1, lab, h="left")
        S.put(r, 2, f"={ref}", fmt="#,##0", color="000000" if d.get("use", True) else "808080")
        for c in range(3, 9):
            S.put(r, c, None)
        r += 1
    tot = r
    S.put(tot, 1, "사업 전체", bold=True, h="left", fill=_CAT)
    S.put(tot, 2, _work_formula(civil_ref, disc_refs, data.get("elec_after_arch", True)),
          bold=True, fmt="#,##0", fill=_CAT)
    S.put(tot, 3, f"={nw_ref}", bold=True, fmt="#,##0", fill=_CAT)
    S.put(tot, 4, per["prep"], bold=True, fill=_CAT)
    S.put(tot, 5, per["wrapup"], bold=True, fill=_CAT)
    S.put(tot, 6, per["commission"], bold=True, fill=_CAT)
    S.put(tot, 7, f"=SUM(B{tot}:F{tot})", bold=True, fmt="#,##0", fill=_CAT)
    S.put(tot, 8, f"=ROUND(G{tot}/365*12,1)", bold=True, fmt="0.0", fill=_CAT)
    r = tot + 2
    notes = [
        "※ 사업 전체 작업일수 = 토목 + 토목 이후 구간(건축·기계·조경은 병행, 전기는 선행 분야 완료 후). "
        "공기 미반영 분야는 계산에서 뺍니다.",
        "※ 비작업일수는 작업일수가 진행되는 본공사 구간의 월별 값을 합산했습니다(별표1).",
        "※ 투입조수(조)는 현장 경험에 따른 가정값이며, 작업일수는 이 가정으로 계산했습니다.",
        *data.get("notes", []),
    ]
    for t in notes:
        S.put(r, 1, t, h="left", border=False, size=8)
        r += 1
    if data.get("start_date"):
        S.put(r + 1, 1, f"착공 예정일: {data['start_date']}", h="left", border=False, size=9)
    _print_setup(summary)

    wb.calculation.fullCalcOnLoad = True       # 받는 쪽 엑셀이 열 때 수식을 모두 계산
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
