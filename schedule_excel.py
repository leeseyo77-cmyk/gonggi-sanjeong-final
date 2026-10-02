# -*- coding: utf-8 -*-
"""
schedule_excel.py — 예정공정표 엑셀 생성

양식은 발주처 제출용 '공사예정표' 샘플(풍각 증설·영해 증설)을 따른다.
  - 왼쪽 두 열: 분야(가설·토목·건축·…·종합시운전) / 세부 구분(토공·구조물공·동 이름 …)
  - 머리행: 기간(개월)+연도 / 경과 개월 1~N / 달력 월. 한 달 = 2열
  - 공종 막대: 초록 선 + 양 끝 동그라미(도형). 공종명은 선 위에 셀 글자로 쓴다
  - 분야 요약: 라벨 상자 + 화살표 + 빨간 요약선, 선행 분야가 끝나는 곳에서 빨간 세로 연결선
  - 아래: 월계(%)·누계(%) 행과 누계 S-커브(파란 점선)
막대를 셀 색칠이 아니라 도형으로 그리는 것은 샘플과 같게 하고, 받는 사람이 엑셀에서 도형을
끌어 일정을 고치기 쉽게 하기 위해서다. openpyxl은 선·타원 같은 임의 도형을 쓰지 못하므로
워크북을 저장한 뒤 DrawingML 파트(xl/drawings/drawing1.xml)를 직접 넣는다.
"""

from __future__ import annotations

import io
import math
import re
import zipfile
from bisect import bisect_right
from datetime import date
from typing import Any, Dict, List, Optional, Sequence
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# 분야 표기 순서(샘플과 같음). 목록에 없는 분야는 뒤에 나온 순서대로 붙인다.
SECTION_ORDER = ["가설공사", "토목공사", "건축공사", "건축기계설비", "기계공사",
                 "전기 및 계측제어공사", "조경공사", "종합시운전"]
# 요약선(라벨 상자·화살표·빨간 선)을 그리지 않는 분야 — 샘플도 가설공사엔 없다
_NO_SUMMARY = {"가설공사", "종합시운전"}
# 월계·누계 작업량에서 빼는 분야(공사준비·정리·시운전은 시공 물량이 아니다)
_NO_LOAD = {"가설공사", "종합시운전"}

MONTH_W = 2                   # 한 달 = 2열(샘플과 같음)
FIRST_MC = 3                  # 첫 달이 시작하는 열(C)
W_A, W_B, W_MONTH, W_NOTE = 12.4, 19.5, 3.6, 9.0
H_TITLE, H_HEAD, H_LANE, H_SUMMARY, H_FOOT = 30.0, 16.0, 21.0, 25.0, 16.0
FONT = "맑은 고딕"

_GREEN, _RED, _NAVY, _BLUE = "00B050", "FF0000", "1F4E79", "6A9BD1"
EMU_PX, EMU_PT = 9525, 12700

_thin = Side(style="thin", color="808080")
_dot = Side(style="dotted", color="A6A6A6")
_med = Side(style="medium", color="404040")
_HDR_FILL = PatternFill("solid", fgColor="F2F2F2")


# ── 입력 정리 ────────────────────────────────────────────────────────────
def _num(v) -> int:
    """None·NaN·문자열을 안전하게 정수로(편집표에서 빈칸이 들어올 수 있다)."""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return 0
    return 0 if math.isnan(f) else int(round(f))


def _txt(v) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    return str(v).strip()


def normalize_rows(rows: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """편집표 행 → 그리기용 행. 시작·기간이 없거나 공종명이 빈 행은 버린다."""
    out = []
    for r in rows:
        s, d, name = _num(r.get("시작(개월차)")), _num(r.get("기간(개월)")), _txt(r.get("공종"))
        if s < 1 or d < 1 or not name:
            continue
        out.append({"분야": _txt(r.get("분야")) or "토목공사", "구분": _txt(r.get("구분")),
                    "공종": name, "조수": max(0, _num(r.get("조수"))),
                    "작업일수": max(0, _num(r.get("작업일수"))), "시작": s, "기간": d})
    return out


def short_label(group: str, name: str) -> str:
    """'구조물공사 (오존접촉조)'처럼 구분명이 앞에 붙은 공종은 괄호 안만 남긴다."""
    m = re.match(r"^\s*(.+?)\s*\((.+)\)\s*$", name or "")
    if m and group and m.group(1).strip() == group.strip():
        return m.group(2).strip()
    return name


def bar_label(row: Dict[str, Any]) -> str:
    """막대 위 글자: 공종명_N조 D일(조수·일수가 있을 때). 조수는 가정값이라 드러내 둔다."""
    lab = short_label(row["구분"], row["공종"])
    if row["조수"] > 0:
        lab += f"_{row['조수']}조"
    if row["작업일수"] > 0:
        lab += f" {row['작업일수']}일"
    return lab


def monthly_progress(rows: Sequence[Dict[str, Any]], n_months: int):
    """월계·누계(%) — 작업량(작업일수×조수, 인·일 대용) 기준으로 막대 기간에 고르게 나눈다."""
    w = [0.0] * n_months
    for r in rows:
        if r["분야"] in _NO_LOAD:
            continue
        load = r["작업일수"] * max(1, r["조수"])
        if load <= 0:
            continue
        for m in range(r["시작"] - 1, min(n_months, r["시작"] - 1 + r["기간"])):
            w[m] += load / r["기간"]
    tot = sum(w)
    if tot <= 0:
        return [0.0] * n_months, [0.0] * n_months
    monthly = [round(x / tot * 100, 1) for x in w]
    cum, acc = [], 0.0
    for x in w:
        acc += x
        cum.append(round(acc / tot * 100, 1))
    cum[-1] = 100.0
    return monthly, cum


def _add_months(d: date, n: int):
    t = d.year * 12 + d.month - 1 + n
    return t // 12, t % 12 + 1


def _col_px(width: float, mdw: int = 7) -> int:
    """엑셀 열 너비(문자 수) → 픽셀. 기본 글꼴(Calibri 11) 최대 숫자폭 7px 기준."""
    return int(((256 * width + int(128 / mdw)) / 256) * mdw)


# ── 도형(DrawingML) ──────────────────────────────────────────────────────
class _Drawing:
    """절대 좌표(EMU)로 받은 도형을 셀 기준 twoCellAnchor로 바꿔 쌓는다.

    셀 기준이라 받는 쪽 엑셀의 기본 글꼴이 달라 열 폭이 조금 변해도 도형이 셀을 따라간다.
    """

    def __init__(self, col_edges: List[int], row_edges: List[int]):
        self.cx, self.ry = col_edges, row_edges
        self.parts: List[str] = []
        self._n = 1

    def _id(self) -> int:
        self._n += 1
        return self._n

    @staticmethod
    def _cell(edges: List[int], v: float):
        i = max(0, min(bisect_right(edges, v) - 1, len(edges) - 2))
        return i, max(0, int(round(v - edges[i])))

    def _anchor(self, x1, y1, x2, y2, inner: str) -> None:
        c1, o1 = self._cell(self.cx, x1)
        r1, p1 = self._cell(self.ry, y1)
        c2, o2 = self._cell(self.cx, x2)
        r2, p2 = self._cell(self.ry, y2)
        self.parts.append(
            "<xdr:twoCellAnchor>"
            f"<xdr:from><xdr:col>{c1}</xdr:col><xdr:colOff>{o1}</xdr:colOff>"
            f"<xdr:row>{r1}</xdr:row><xdr:rowOff>{p1}</xdr:rowOff></xdr:from>"
            f"<xdr:to><xdr:col>{c2}</xdr:col><xdr:colOff>{o2}</xdr:colOff>"
            f"<xdr:row>{r2}</xdr:row><xdr:rowOff>{p2}</xdr:rowOff></xdr:to>"
            f"{inner}<xdr:clientData/></xdr:twoCellAnchor>")

    @staticmethod
    def _xfrm(x, y, w, h, flip_v=False) -> str:
        fv = ' flipV="1"' if flip_v else ""
        return (f'<a:xfrm{fv}><a:off x="{int(x)}" y="{int(y)}"/>'
                f'<a:ext cx="{max(0, int(w))}" cy="{max(0, int(h))}"/></a:xfrm>')

    def line(self, x1, y1, x2, y2, color: str, w_emu: int, dash: Optional[str] = None) -> None:
        if x2 < x1:
            x1, y1, x2, y2 = x2, y2, x1, y1
        flip_v = y2 < y1
        top, bot = (y2, y1) if flip_v else (y1, y2)
        i = self._id()
        dash_xml = f'<a:prstDash val="{dash}"/>' if dash else ""
        inner = (f'<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="{i}" name="선 {i}"/>'
                 "<xdr:cNvCxnSpPr/></xdr:nvCxnSpPr><xdr:spPr>"
                 f"{self._xfrm(x1, top, x2 - x1, bot - top, flip_v)}"
                 '<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
                 f'<a:ln w="{w_emu}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
                 f"{dash_xml}</a:ln></xdr:spPr></xdr:cxnSp>")
        self._anchor(x1, top, x2, bot, inner)

    def shape(self, prst: str, x1, y1, x2, y2, fill: str, line_color: str, line_w: int = 9525,
              text: Optional[str] = None, text_color: str = "000000", size: int = 800,
              bold: bool = False) -> None:
        i = self._id()
        tx = ""
        if text is not None:
            tx = ('<xdr:txBody><a:bodyPr vertOverflow="overflow" horzOverflow="overflow" wrap="none" '
                  'lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="ctr"/><a:lstStyle/>'
                  '<a:p><a:pPr algn="ctr"/><a:r>'
                  f'<a:rPr lang="ko-KR" altLang="en-US" sz="{size}" b="{1 if bold else 0}">'
                  f'<a:solidFill><a:srgbClr val="{text_color}"/></a:solidFill>'
                  f'<a:latin typeface="{FONT}"/><a:ea typeface="{FONT}"/></a:rPr>'
                  f"<a:t>{escape(text)}</a:t></a:r></a:p></xdr:txBody>")
        inner = (f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{i}" name="도형 {i}"/>'
                 "<xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr>"
                 f"{self._xfrm(x1, y1, x2 - x1, y2 - y1)}"
                 f'<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>'
                 f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>'
                 f'<a:ln w="{line_w}"><a:solidFill><a:srgbClr val="{line_color}"/></a:solidFill></a:ln>'
                 f"</xdr:spPr>{tx}</xdr:sp>")
        self._anchor(x1, y1, x2, y2, inner)

    def xml(self) -> str:
        return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                + "".join(self.parts) + "</xdr:wsDr>")


_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_DRAW_REL = (f'<Relationship Id="rIdSchedDrawing" Type="{_REL_NS}/drawing" '
             'Target="../drawings/drawing1.xml"/>')


def _inject_drawing(xlsx: bytes, drawing_xml: str) -> bytes:
    """openpyxl이 저장한 xlsx의 첫 시트에 도형 파트를 붙인다."""
    sheet, rels = "xl/worksheets/sheet1.xml", "xl/worksheets/_rels/sheet1.xml.rels"
    zin = zipfile.ZipFile(io.BytesIO(xlsx))
    names = zin.namelist()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            data = zin.read(n)
            if n == sheet:
                s = data.decode("utf-8")
                tag = f'<drawing xmlns:r="{_REL_NS}" r:id="rIdSchedDrawing"/>'
                # 스키마 순서상 drawing은 아래 요소들보다 앞이어야 한다
                cut = [i for i in (s.find(t) for t in ("<legacyDrawing", "<picture", "<oleObjects",
                                                       "<controls", "<webPublishItems",
                                                       "<tableParts", "<extLst")) if i >= 0]
                at = min(cut) if cut else s.rfind("</worksheet>")
                data = (s[:at] + tag + s[at:]).encode("utf-8")
            elif n == "[Content_Types].xml":
                data = data.decode("utf-8").replace(
                    "</Types>", '<Override PartName="/xl/drawings/drawing1.xml" ContentType='
                                '"application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>'
                ).encode("utf-8")
            elif n == rels:
                data = data.decode("utf-8").replace(
                    "</Relationships>", _DRAW_REL + "</Relationships>").encode("utf-8")
            zout.writestr(n, data)
        if rels not in names:
            zout.writestr(rels, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                                '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
                                f'2006/relationships">{_DRAW_REL}</Relationships>')
        zout.writestr("xl/drawings/drawing1.xml", drawing_xml.encode("utf-8"))
    return out.getvalue()


# ── 본체 ────────────────────────────────────────────────────────────────
def _sections(rows: List[Dict[str, Any]]):
    """[(분야, [(구분, [행…])…])…] — 분야는 SECTION_ORDER 순, 구분은 나온 순서."""
    order = {s: i for i, s in enumerate(SECTION_ORDER)}
    seen: List[str] = []
    for r in rows:
        if r["분야"] not in seen:
            seen.append(r["분야"])
    first = {s: i for i, s in enumerate(seen)}
    seen = sorted(seen, key=lambda s: (order.get(s, len(order)), first[s]))
    out = []
    for sec in seen:
        groups: Dict[str, List[Dict[str, Any]]] = {}
        for r in rows:
            if r["분야"] == sec:
                groups.setdefault(r["구분"], []).append(r)
        out.append((sec, list(groups.items())))
    return out


def build_schedule_xlsx(rows: Sequence[Dict[str, Any]], start_date: date, title: str,
                        notes: Sequence[str] = ()) -> bytes:
    """편집표 행으로 샘플 양식의 예정공정표 엑셀(바이트)을 만든다.

    rows: 분야·구분·공종·조수·시작(개월차)·기간(개월)·작업일수 키를 가진 dict 목록
    """
    data = normalize_rows(rows)
    if not data:
        raise ValueError("그릴 공정이 없습니다(시작·기간·공종명을 확인하세요).")
    n_months = max(r["시작"] + r["기간"] - 1 for r in data)
    last_mc = FIRST_MC + n_months * MONTH_W - 1      # 마지막 달의 마지막 열
    note_col = last_mc + 1

    wb = Workbook()
    ws = wb.active
    ws.title = "예정공정표"

    # 열 폭
    ws.column_dimensions["A"].width = W_A
    ws.column_dimensions["B"].width = W_B
    for c in range(FIRST_MC, last_mc + 1):
        ws.column_dimensions[get_column_letter(c)].width = W_MONTH
    ws.column_dimensions[get_column_letter(note_col)].width = W_NOTE

    def cell(r, c, v=None, size=9, bold=False, color="000000", h="center", v_al="center",
             wrap=False, fill=None, border=None):
        x = ws.cell(row=r, column=c, value=v)
        x.font = Font(name=FONT, size=size, bold=bold, color=color)
        x.alignment = Alignment(horizontal=h, vertical=v_al, wrap_text=wrap)
        if fill:
            x.fill = fill
        if border:
            x.border = border
        return x

    def box(r1, c1, r2, c2, side=_thin):
        """병합 영역 바깥 테두리."""
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                b = ws.cell(row=rr, column=cc).border
                ws.cell(row=rr, column=cc).border = Border(
                    left=side if cc == c1 else b.left, right=side if cc == c2 else b.right,
                    top=side if rr == r1 else b.top, bottom=side if rr == r2 else b.bottom)

    # ── 머리행(1~4행) ──
    ws.row_dimensions[1].height = H_TITLE
    for r in (2, 3, 4):
        ws.row_dimensions[r].height = H_HEAD
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=note_col)
    cell(1, 1, title, size=16, bold=True)
    ws.merge_cells("A2:B2")
    cell(2, 1, "기간(개월)", bold=True, fill=_HDR_FILL)
    ws.merge_cells("A3:B4")
    cell(3, 1, "공종", bold=True, fill=_HDR_FILL)
    ws.merge_cells(start_row=2, start_column=note_col, end_row=4, end_column=note_col)
    cell(2, note_col, "비고", bold=True, fill=_HDR_FILL)
    years: Dict[int, List[int]] = {}
    for m in range(n_months):
        y, mm = _add_months(start_date, m)
        c = FIRST_MC + m * MONTH_W
        years.setdefault(y, []).append(c)
        for row_i, val in ((3, m + 1), (4, mm)):
            ws.merge_cells(start_row=row_i, start_column=c, end_row=row_i, end_column=c + MONTH_W - 1)
            cell(row_i, c, val, bold=True, fill=_HDR_FILL)
            box(row_i, c, row_i, c + MONTH_W - 1)
    for y, cols in years.items():
        c1, c2 = min(cols), max(cols) + MONTH_W - 1
        ws.merge_cells(start_row=2, start_column=c1, end_row=2, end_column=c2)
        cell(2, c1, f"{y}년", bold=True, fill=_HDR_FILL)
        box(2, c1, 2, c2)
    box(2, 1, 2, 2)
    box(3, 1, 4, 2)
    box(2, note_col, 4, note_col)

    # ── 본문: 분야 → (요약줄) → 구분 → 공종 줄 ──
    lanes: List[Dict[str, Any]] = []
    r = 5
    sec_blocks = []
    for sec, groups in _sections(data):
        sec_r1 = r
        rows_in = [x for _, g in groups for x in g]
        s_start = min(x["시작"] for x in rows_in)
        s_end = max(x["시작"] + x["기간"] - 1 for x in rows_in)
        if sec not in _NO_SUMMARY:
            ws.row_dimensions[r].height = H_SUMMARY
            lanes.append({"kind": "summary", "row": r, "sec": sec, "start": s_start, "end": s_end,
                          "groups": groups})
            r += 1
        for grp, items in groups:
            g_r1 = r
            for it in items:
                ws.row_dimensions[r].height = H_LANE
                lanes.append({"kind": "act", "row": r, "sec": sec, "item": it})
                r += 1
            ws.merge_cells(start_row=g_r1, start_column=2, end_row=r - 1, end_column=2)
            cell(g_r1, 2, grp, size=9, wrap=True)
            box(g_r1, 2, r - 1, 2)
            box(g_r1, FIRST_MC, r - 1, last_mc)
        ws.merge_cells(start_row=sec_r1, start_column=1, end_row=r - 1, end_column=1)
        cell(sec_r1, 1, sec, size=9, bold=True, wrap=True)
        sec_blocks.append((sec_r1, r - 1))
    body_r1, body_r2 = 5, r - 1

    # 달 경계 점선 + 분야 경계 굵은 선 + 비고 열
    for rr in range(body_r1, body_r2 + 1):
        for m in range(n_months):
            c = FIRST_MC + m * MONTH_W
            b = ws.cell(row=rr, column=c).border
            ws.cell(row=rr, column=c).border = Border(left=_dot if m else _thin, right=b.right,
                                                      top=b.top, bottom=b.bottom)
        b = ws.cell(row=rr, column=note_col).border
        ws.cell(row=rr, column=note_col).border = Border(left=_thin, right=_thin, top=b.top,
                                                         bottom=b.bottom)
    for s1, s2 in sec_blocks:
        box(s1, 1, s2, note_col, side=_med)

    # 공종 글자(선 위) · 요약선 글자
    for ln in lanes:
        if ln["kind"] == "act":
            it = ln["item"]
            c = FIRST_MC + (it["시작"] - 1) * MONTH_W
            color = _RED if ln["sec"] == "종합시운전" else _GREEN
            cell(ln["row"], c, bar_label(it), size=8, color=color, h="left", v_al="top")
        else:
            mid = (ln["start"] - 1 + ln["end"]) / 2 + 0.6     # 상자·화살표 뒤쪽 가운데
            c = FIRST_MC + min(n_months - 1, int(mid)) * MONTH_W
            cell(ln["row"], c, ln["sec"], size=11, bold=True, color=_RED, h="center", v_al="top")

    # ── 월계·누계 ──
    monthly, cum = monthly_progress(data, n_months)
    for lab, vals in (("월계(%)", monthly), ("누계(%)", cum)):
        ws.row_dimensions[r].height = H_FOOT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        cell(r, 1, lab, bold=True, fill=_HDR_FILL)
        box(r, 1, r, 2)
        for m, v in enumerate(vals):
            c = FIRST_MC + m * MONTH_W
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + MONTH_W - 1)
            x = cell(r, c, v, size=8)
            x.number_format = "0.0"
            box(r, c, r, c + MONTH_W - 1)
        box(r, note_col, r, note_col)
        r += 1
    note_lines = ["주) 본 예정공정표는 전체 공정계획용이며, 시공자는 착공 시 자원투입계획을 반영한 "
                  "세부 예정공정표를 작성하여 발주처(건설사업관리자)의 승인을 받아야 함",
                  "주) 월계·누계는 공종별 작업량(작업일수×투입조수) 기준이며 공사비 보할과 다를 수 있음",
                  *notes]
    for t in note_lines:
        ws.row_dimensions[r].height = H_FOOT
        cell(r, 1, t, size=8, h="left")
        r += 1

    # ── 인쇄 설정: A3 가로 한 장 너비 ──
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 8
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_area = f"A1:{get_column_letter(note_col)}{r - 1}"
    ws.freeze_panes = "C5"

    # ── 도형 좌표계(EMU) ──
    widths = {1: W_A, 2: W_B, note_col: W_NOTE}
    col_edges, acc = [0], 0
    for c in range(1, note_col + 40):
        w = widths.get(c, W_MONTH if FIRST_MC <= c <= last_mc else 8.43)
        acc += _col_px(w) * EMU_PX
        col_edges.append(acc)
    row_edges, acc = [0], 0
    for rr in range(1, r + 40):
        h = ws.row_dimensions[rr].height or 15.0
        acc += int(h * EMU_PT)
        row_edges.append(acc)
    month_emu = _col_px(W_MONTH) * EMU_PX * MONTH_W

    def mx(mf: float) -> float:
        """0부터 센 개월(실수) → x(EMU)."""
        return col_edges[FIRST_MC - 1] + mf * month_emu

    def row_y(row: int, frac: float) -> float:
        return row_edges[row - 1] + frac * (row_edges[row] - row_edges[row - 1])

    dr = _Drawing(col_edges, row_edges)
    rad = 3.5 * EMU_PX
    summary_y: Dict[str, float] = {}
    sec_span: Dict[str, tuple] = {}
    for ln in lanes:
        if ln["kind"] == "act":
            it = ln["item"]
            x1, x2 = mx(it["시작"] - 1), mx(it["시작"] - 1 + it["기간"])
            y = row_y(ln["row"], 0.74)
            color = _RED if ln["sec"] == "종합시운전" else _GREEN
            dr.line(x1, y, x2, y, color, 15875)
            for xx in (x1, x2):
                # 첫 달 왼쪽 경계를 넘어 B열에 걸리면 엑셀이 B열 폭으로 늘여 그린다 → 안쪽으로 민다
                xx = max(xx, mx(0) + rad)
                dr.shape("ellipse", xx - rad, y - rad, xx + rad, y + rad, "FFFFFF", _RED, 9525)
        else:
            row = ln["row"]
            y = row_y(row, 0.62)
            x1, x2 = mx(ln["start"] - 1), mx(ln["end"])
            summary_y[ln["sec"]] = y
            sec_span[ln["sec"]] = (ln["start"], ln["end"])
            # 라벨 상자: 착수 시점 바로 앞(샘플처럼), 첫 열을 넘지 않게
            bw = 1.8 * month_emu
            if ln["start"] > 1:
                bx1, bx2 = max(mx(0), x1 - bw), x1
            else:
                bx1, bx2 = x1, x1 + bw
            top, bot = row_y(row, 0.06), row_y(row, 0.94)
            mid = (top + bot) / 2
            sub = (f"{len(ln['groups'])}개 동" if ln["sec"].startswith("건축") and len(ln["groups"]) > 1
                   else (ln["groups"][0][0] if len(ln["groups"]) == 1 and ln["groups"][0][0] else "주공정"))
            dr.shape("rect", bx1, top, bx2, mid, "FFC000", "000000", 6350, ln["sec"], size=800, bold=True)
            dr.shape("rect", bx1, mid, bx2, bot, "FFF2CC", "000000", 6350, sub, size=700)
            # 화살표 + 빨간 요약선
            ax1 = max(x1, bx2) + 0.05 * month_emu
            ax2 = min(x2, ax1 + 0.7 * month_emu)
            if ax2 > ax1:
                dr.shape("rightArrow", ax1, y - 4.5 * EMU_PT, ax2, y + 4.5 * EMU_PT, _RED, _NAVY, 12700)
            if x2 > ax2:
                dr.line(ax2, y, x2, y, _RED, 25400)

    # 분야 연결: 선행 분야가 끝나는 지점에서 다음 분야 요약선으로 내려오는 빨간 세로선
    secs = list(summary_y)
    for i, s in enumerate(secs[1:], start=1):
        st_m = sec_span[s][0]
        prev = [p for p in secs[:i] if sec_span[p][1] <= st_m]
        if prev:
            p = max(prev, key=lambda q: sec_span[q][1])
            x = mx(st_m - 1)
            dr.line(x, summary_y[p], x, summary_y[s], _RED, 19050)

    # 누계 S-커브: 본문 아래가 0%, 위가 100%
    top_y, bot_y = row_edges[body_r1 - 1], row_edges[body_r2]
    pts = [(mx(0), bot_y)] + [(mx(m + 1), bot_y - cum[m] / 100.0 * (bot_y - top_y))
                               for m in range(n_months)]
    if any(cum):
        for (xa, ya), (xb, yb) in zip(pts, pts[1:]):
            dr.line(xa, ya, xb, yb, _BLUE, 34925, dash="dash")

    buf = io.BytesIO()
    wb.save(buf)
    return _inject_drawing(buf.getvalue(), dr.xml())
