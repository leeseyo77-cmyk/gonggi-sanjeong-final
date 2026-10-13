"""노무비 역산 — 직종별 인수는 병목 직종 기준(1조 = 직종별 1인)으로 환산하는지 확인한다.

품셈의 직종별 인수는 여러 직종이 함께 드는 조합이라, 인수를 모두 더해 나누면(1인 기준)
작업일수가 1.5~3배 길게 나왔다. 풍각 샘플의 호표로 실제 값을 대조한다(샘플 없으면 건너뜀).
실행: python -m unittest discover -s tests -v
"""
import unittest

import openpyxl

from _support import SAMPLES, have

import universal_parser as up

PUNGGAK = SAMPLES / "풍각 공공하수처리장 증설사업260609.xlsx"


@unittest.skipUnless(have(PUNGGAK), "풍각 샘플 없음")
class BottleneckTradeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        wb = openpyxl.load_workbook(PUNGGAK, data_only=True, keep_links=False)
        _, traces = up.parse_labor_derived_by_template(wb, up.detect_template(wb), with_trace=True)
        cls.by_ho = {t.get("호표"): t for t in traces.values()}

    def test_two_trades(self):
        # 제96호표 보강토옹벽 블록: 특별인부 0.21 + 보통인부 0.09 → 병목 특별인부
        t = self.by_ho[96]
        self.assertEqual(t["방식"], "직종별 인수(병목 직종)")
        self.assertEqual(t["병목 직종"], "특별인부")
        self.assertAlmostEqual(t["일작업량"], 1 / 0.21, places=3)
        self.assertAlmostEqual(t["1인 기준 일작업량"], 1 / 0.30, places=3)

    def test_four_trades(self):
        # 제100호표 세라믹 방수: 미장공 0.04 + 도장공 0.03 + 보통인부 0.03 + 특별인부 0.02
        t = self.by_ho[100]
        self.assertEqual(t["병목 직종"], "미장공")
        self.assertAlmostEqual(t["일작업량"], 25.0, places=3)        # 1인 기준이면 8.33
        self.assertAlmostEqual(t["인수합"], 0.12, places=6)


def _sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    for r in rows:
        ws.append(r)


def _row(name, spec=None, qty=None, unit=None, labor=None, ref=None):
    r = [name, spec, qty, unit, None, None, None, labor, None, None, None, None, ref]
    return r


class SangeunAQTest(unittest.TestCase):
    """일위대가 → 산근 A·Q 따라가기(샘플 불필요). 표준형: 구성 행 비고 열에 '산근 N호표'."""

    @classmethod
    def setUpClass(cls):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        _sheet(wb, "노무비목록표", [[None, "보통인부", None, None, 172068]])
        _sheet(wb, "일위대가표", [
            ["제 1 호표"], _row("띠장 설치 및 철거", "L=9∼11m", None, "m"),
            _row("H-Beam 설치/띠장", "L=9∼11m", 1, "m", 25191, "산근   4호표"),
            _row("H-Beam 해체/띠장", "L=9∼11m", 1, "m", 13297, "산근  10호표"),
            _row("합  계", None, None, None, 38488),
            ["제 2 호표"], _row("PVC PIPE설치", "D=100", None, "개소"),
            _row("PVC PIPE 접합", "D=100", 1, "개소", 9000, "산근 5호표"),
            _row("보통인부", None, 0.04, "인", 6882),
            _row("합  계", None, None, None, 15882),
            ["제 3 호표"], _row("플랫타이 구멍메우기", None, None, "개소"),
            _row("보통인부", None, 0.62, "인", 106682),
            _row("1개소당", "1÷100개소", 1, "%", 1066),
            _row("합  계", None, None, None, 1066),
            ["제 4 호표"], _row("단위 다른 산근", None, None, "m"),
            _row("어떤 작업", None, 1, "TON", 5000, "산근 6호표"),
            _row("보통인부", None, 0.5, "인", 86034),
            _row("합  계", None, None, None, 91034),
        ])
        code_daily = {4: (80.0, "m"), 10: (110.0, "m"), 5: (25.0, "개소"), 6: (100.0, "㎥")}
        cls.res, cls.tr = up.parse_labor_derived_by_template(wb, up.TEMPLATES[0], with_trace=True,
                                                             code_daily=code_daily)

    def test_install_and_removal_days_add(self):
        # 설치 A=80, 해체 A=110 → 1m당 1/80 + 1/110일 → 46.3m/일 (A값을 더한 190이 아님)
        d, u, job = self.res[1]
        self.assertEqual(job, "산근 A·Q")
        self.assertAlmostEqual(d, 1 / (1 / 80 + 1 / 110), places=3)
        self.assertEqual(self.tr[1]["방식"], "산근 A·Q(일위대가 경유)")
        self.assertEqual([p["산근"] for p in self.tr[1]["산근 작업"]], [4, 10])

    def test_rest_by_trade_manhours(self):
        # A·Q가 없는 나머지(보통인부 0.04인)는 직종별 인수로 일수를 더한다
        self.assertAlmostEqual(self.res[2][0], 1 / (1 / 25 + 0.04), places=3)

    def test_per_100_base(self):
        # '1개소당 | 1÷100개소' — 0.62인은 100개소 기준
        self.assertAlmostEqual(self.res[3][0], 100 / 0.62, places=2)
        self.assertEqual(self.tr[3]["방식"], "직종별 인수(병목 직종)")

    def test_unit_mismatch_falls_back(self):
        # 산근 단위(㎥)와 구성 단위(TON)가 다르면 A·Q를 쓰지 않고 인수로
        self.assertEqual(self.tr[4]["방식"], "직종별 인수(병목 직종)")
        self.assertAlmostEqual(self.res[4][0], 2.0, places=4)

    def test_a_value_preferred_over_rounded_q(self):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        _sheet(wb, "단가산출근거", [
            [None, "제   7호표"], [None, "콘크리트 타설 / ㎥"],
            [None, " -일당 시공량: A=63 ㎥/일 "],
            [None, " -시간당 시공량: Q=A/8 hr = 7.88 ㎥/hr "],
            # 할증계수를 곱한 최종값(광주 제32호표)
            [None, "제  32호표"], [None, "소파보수(표층) / ㎡"],
            [None, " -일당 시공량: A=400 ㎡/일 *1.26 (할증계수) = 504.00"],
            [None, " -시간당 시공량: Q=A/8 hr = 63.00 ㎡/hr "],
            # A가 일부 작업분(인력 50%)일 때는 장비 Q를 쓴다(본대 제140호표)
            [None, "제 140호표"], [None, "터파기 굴착기 50%+인력 50% / ㎥"],
            [None, "Q=3600*qo*K*f*E/Cm= 6.22 ㎥/hr"],
            [None, " -일당 시공량: A=3.6 ㎥/일"],
            [None, " -시간당 시공량: Q=A/8 hr = 0.45 ㎥/hr"],
        ])
        got = up.parse_unit_price_generic(wb, up.TEMPLATES[0])
        self.assertEqual(got[7], (63.0, "㎥"))
        self.assertEqual(got[32], (504.0, "㎡"))
        self.assertAlmostEqual(got[140][0], 6.22 * 8, places=3)


if __name__ == "__main__":
    unittest.main()
