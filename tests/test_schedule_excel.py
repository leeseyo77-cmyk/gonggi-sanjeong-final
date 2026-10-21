"""schedule_excel 단위 테스트 — 합성 공정으로 예정공정표 엑셀을 만들어 구조를 확인한다(수 초).

실행: python -m unittest discover -s tests -v
"""
import io
import unittest
import zipfile
from datetime import date
from xml.dom.minidom import parseString

import openpyxl

import _support  # noqa: F401  (저장소 경로 등록)
import schedule_excel as sx

ROWS = [
    {"분야": "가설공사", "구분": "공사준비", "공종": "공사준비(인허가)", "조수": 0,
     "시작(개월차)": 1, "기간(개월)": 2, "작업일수": 60},
    {"분야": "토목공사", "구분": "토공", "공종": "토공", "조수": 2,
     "시작(개월차)": 3, "기간(개월)": 4, "작업일수": 90},
    {"분야": "토목공사", "구분": "구조물공사", "공종": "구조물공사 (오존접촉조)", "조수": 3,
     "시작(개월차)": 3, "기간(개월)": 10, "작업일수": 280},
    {"분야": "건축공사", "구분": "A동", "공종": "철근콘크리트공사", "조수": 4,
     "시작(개월차)": 13, "기간(개월)": 5, "작업일수": 120},
    {"분야": "건축공사", "구분": "B동", "공종": "철근콘크리트공사", "조수": 2,
     "시작(개월차)": 13, "기간(개월)": 4, "작업일수": 100},
    {"분야": "종합시운전", "구분": "시운전", "공종": "시운전/시설인계", "조수": 0,
     "시작(개월차)": 18, "기간(개월)": 2, "작업일수": 60},
]


class BuildTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.xlsx = sx.build_schedule_xlsx(ROWS, date(2026, 1, 1), "테스트 예정공정표",
                                          notes=["※ 막대의 'N조'는 가정값"], flood_months=[9, 10])
        cls.zip = zipfile.ZipFile(io.BytesIO(cls.xlsx))
        cls.ws = openpyxl.load_workbook(io.BytesIO(cls.xlsx)).active

    def test_drawing_part_is_wired(self):
        names = self.zip.namelist()
        self.assertIn("xl/drawings/drawing1.xml", names)
        parseString(self.zip.read("xl/drawings/drawing1.xml"))          # 올바른 XML
        self.assertIn(b"rIdSchedDrawing", self.zip.read("xl/worksheets/sheet1.xml"))
        self.assertIn(b"rIdSchedDrawing", self.zip.read("xl/worksheets/_rels/sheet1.xml.rels"))
        self.assertIn(b"/xl/drawings/drawing1.xml", self.zip.read("[Content_Types].xml"))

    def test_shapes_drawn(self):
        xml = self.zip.read("xl/drawings/drawing1.xml").decode("utf-8")
        self.assertEqual(xml.count('prst="ellipse"'), 2 * len(ROWS))   # 공종마다 양 끝 동그라미
        self.assertEqual(xml.count('prst="rightArrow"'), 2)             # 토목·건축 요약 화살표
        self.assertIn('prstDash val="dash"', xml)                       # S-커브

    def test_sections_in_sample_order(self):
        secs = [self.ws.cell(r, 1).value for r in range(5, self.ws.max_row + 1)
                if self.ws.cell(r, 1).value]
        self.assertEqual(secs[:4], ["가설공사", "토목공사", "건축공사", "종합시운전"])

    def test_header_and_month_count(self):
        self.assertEqual(self.ws.cell(1, 1).value, "테스트 예정공정표")
        self.assertEqual(self.ws.cell(3, 3).value, 1)                   # 경과 1개월
        self.assertEqual(self.ws.cell(2, 3).value, "2026년")
        self.assertEqual(self.ws.max_column, 3 + 19 * 2)               # 19개월 × 2열 + 비고

    def test_progress_rows(self):
        vals = {self.ws.cell(r, 1).value: r for r in range(5, self.ws.max_row + 1)}
        cum_row = vals["누계(%)"]
        cum = [self.ws.cell(cum_row, c).value for c in range(3, 3 + 19 * 2, 2)]
        self.assertEqual(cum[-1], 100)
        self.assertEqual(cum, sorted(cum))
        self.assertEqual(cum[0], 0)                                     # 공사준비는 작업량 0

    def test_flood_season_row_and_shading(self):
        # 풍각 샘플: 수방기간(9·10월) 열 음영 + 월계 위 '수방기간' 행
        vals = {self.ws.cell(r, 1).value: r for r in range(5, self.ws.max_row + 1)}
        fr = vals["수방기간"]
        self.assertEqual(fr + 1, vals["월계(%)"])
        sep = 3 + 8 * 2                                                 # 2026-09 = 9번째 달
        self.assertEqual(self.ws.cell(fr, sep).value, "수방기간")
        self.assertEqual(self.ws.cell(6, sep).fill.fgColor.rgb[-6:], "DDEBF7")
        self.assertNotEqual((self.ws.cell(6, sep - 2).fill.fgColor.rgb or "")[-6:], "DDEBF7")

    def test_notes_written(self):
        texts = [self.ws.cell(r, 1).value for r in range(1, self.ws.max_row + 1)]
        self.assertIn("※ 막대의 'N조'는 가정값", texts)


class HelperTest(unittest.TestCase):
    def test_short_label(self):
        self.assertEqual(sx.short_label("구조물공사", "구조물공사 (오존접촉조)"), "오존접촉조")
        self.assertEqual(sx.short_label("토공", "토공"), "토공")
        self.assertEqual(sx.short_label("A동", "철근콘크리트공사"), "철근콘크리트공사")

    def test_bar_label_is_name_only(self):
        # 샘플처럼 막대 위에는 공종명만 — 조수 가정은 아래 주석으로
        row = sx.normalize_rows([ROWS[3]])[0]
        self.assertEqual(sx.bar_label(row), "철근콘크리트공사")

    def test_fractional_months(self):
        r = sx.normalize_rows([{"공종": "터파기", "시작(개월차)": 3.5, "기간(개월)": 1.25}])[0]
        self.assertEqual((r["s0"], r["e0"]), (2.5, 3.75))

    def test_sequential_bars_share_a_lane(self):
        rows = sx.normalize_rows([
            {"구분": "토공", "공종": "표토제거", "시작(개월차)": 1, "기간(개월)": 1},
            {"구분": "토공", "공종": "터파기", "시작(개월차)": 3, "기간(개월)": 2},
            {"구분": "토공", "공종": "되메우기", "시작(개월차)": 3.5, "기간(개월)": 1}])   # 터파기와 겹침
        lanes = sx.pack_lanes(rows)
        self.assertEqual([[x["공종"] for x in ln] for ln in lanes], [["표토제거", "터파기"], ["되메우기"]])

    def test_progress_is_integer_and_sums_to_100(self):
        rows = sx.normalize_rows(ROWS)
        monthly, cum = sx.monthly_progress(rows, 19)
        self.assertTrue(all(isinstance(v, int) for v in monthly + cum))
        self.assertEqual(sum(monthly), 100)
        self.assertEqual(cum[-1], 100)

    def test_bad_rows_are_skipped(self):
        bad = [{"공종": "x", "시작(개월차)": float("nan"), "기간(개월)": 2},
               {"공종": "", "시작(개월차)": 1, "기간(개월)": 2},
               {"공종": "y", "시작(개월차)": 1, "기간(개월)": 0}]
        self.assertEqual(sx.normalize_rows(bad), [])
        with self.assertRaises(ValueError):
            sx.build_schedule_xlsx(bad, date(2026, 1, 1), "t")

    def test_missing_section_defaults_to_civil(self):
        r = sx.normalize_rows([{"공종": "토공", "시작(개월차)": 1, "기간(개월)": 1}])
        self.assertEqual(r[0]["분야"], "토목공사")


if __name__ == "__main__":
    unittest.main()
